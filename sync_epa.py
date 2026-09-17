"""Actualiza las encuestas embebidas en el panel: python sync_epa.py.

Lee únicamente la hoja principal Export/BaseEpaPersona. Las hojas auxiliares
no se suman porque son ajustes/subconjuntos de la base. Requiere openpyxl.
"""
import argparse
import datetime as dt
import json
import os
import tempfile
from urllib.parse import quote
from pathlib import Path
import re
import warnings
import zipfile
import xml.etree.ElementTree as ET
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
# Mantener los nombres originales cuando todavía no hay equivalencia confirmada.
PDV = {4628: 'MALL CENTER CURICÓ', 4172: 'TALCA MALL',
       4165: 'CURICÓ CAMILO HENRÍQUEZ', 4939: 'TALCA CENTRO',
       4620: 'LINARES INDEPENDENCIA', 4169: 'CHILLÁN EL ROBLE',
       5053: 'MALL ARAUCO CHILLÁN', 4167: 'LOS ÁNGELES LAUTARO',
       4579: 'MALL PLAZA LOS ÁNGELES'}

def value(v):
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return v.isoformat()
    return v

def build(source):
    records, sources, headers = [], [], None
    for path in sorted(source.glob('*.xlsx')):
        if path.name.startswith('~$'):
            continue
        with warnings.catch_warnings():
            warnings.filterwarnings('ignore', message='Cell .* is marked as a date')
            book = load_workbook(path, read_only=True, data_only=True)
            sheet = next((s for s in book if s.title in ('Export', 'BaseEpaPersona')), None)
            if sheet is None:
                raise ValueError(f'{path.name}: falta hoja de encuestas')
            rows = list(sheet.iter_rows(values_only=True))
            current_headers = list(rows[0])
            if headers is None:
                headers = current_headers
            if current_headers != headers:
                raise ValueError(f'{path.name}: cambiaron las columnas; revisar importación')
            # Recuperar valores originales con formatos de fecha inválidos en Excel.
            with zipfile.ZipFile(path) as archive:
                xml = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
                ns = {'x': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                raw = {c.attrib['r']: c.findtext('x:v', namespaces=ns)
                       for c in xml.findall('.//x:c', ns)}
            from openpyxl.utils import get_column_letter
            count = 0
            for number, row in enumerate(rows[1:], 2):
                if not any(v is not None for v in row):
                    continue
                vals = [value(v) for v in row]
                for i, v in enumerate(vals):
                    if v == '#VALUE!':
                        original = raw.get(f'{get_column_letter(i+1)}{number}')
                        if original and original.isdigit():
                            vals[i] = original
                year, month = int(row[5]), int(row[4])
                if not 2000 <= year <= 2100 or not 1 <= month <= 12:
                    raise ValueError(f'{path.name}:{number}: período inválido')
                code = int(row[43]) if row[43] is not None else None
                records.append({'mes': f'{year:04d}-{month:02d}',
                                'sucursal': PDV.get(code, str(row[45] or 'Sin sucursal')),
                                'fuente': len(sources), 'fila': number, 'valores': vals})
                count += 1
            sources.append({'archivo': path.name, 'hoja': sheet.title, 'registros': count})
            book.close()
    if not records:
        raise ValueError('No se encontraron encuestas. El HTML no se modificó.')
    return {'actualizado': dt.datetime.now().isoformat(timespec='seconds'),
            'columnas': headers, 'fuentes': sources, 'registros': records}

def update_html(html_path, data):
    html = html_path.read_text(encoding='utf-8')
    pattern = r'<script id="epa-data" type="application/json">.*?</script>'
    existing = re.search(pattern, html, re.S)
    if existing:
        previous = json.loads(existing.group().split('>', 1)[1].rsplit('</script>', 1)[0])
        if {k:v for k,v in previous.items() if k != 'actualizado'} == {k:v for k,v in data.items() if k != 'actualizado'}:
            print('EPA sin cambios; se conserva la fecha de actualización.')
            return False
    payload = json.dumps(data, ensure_ascii=False, separators=(',', ':'), allow_nan=False).replace('<', '\\u003c').replace('>', '\\u003e').replace('&', '\\u0026')
    block = '<script id="epa-data" type="application/json">' + payload + '</script>'
    if existing:
        html = re.sub(pattern, lambda _: block, html, flags=re.S)
    else:
        if '<script>\n' not in html:
            raise ValueError('No se encontró el bloque de scripts del panel.')
        html = html.replace('<script>\n', block + '\n<script>\n', 1)
    temporary = html_path.with_suffix('.html.tmp')
    temporary.write_text(html, encoding='utf-8')
    temporary.replace(html_path)
    return True

def download_graph(destination):
    import requests
    site = os.environ.get('EPA_SITE_ID') or 'clcomercialexpress.sharepoint.com,d331a44f-8002-4cab-8db4-061cb13eb493,1d32dc11-bde8-40af-8d21-d070452a7b2b'
    folder = os.environ.get('EPA_FOLDER_PATH') or 'BASES ENTEL/EPA_CALIDAD'
    auth = requests.post(f'https://login.microsoftonline.com/{quote(os.environ["TENANT_ID"], safe="")}/oauth2/v2.0/token',
        data={'client_id': os.environ['CLIENT_ID'], 'client_secret': os.environ['CLIENT_SECRET'],
              'scope': 'https://graph.microsoft.com/.default', 'grant_type': 'client_credentials'}, timeout=40)
    if auth.status_code != 200:
        raise ValueError(f'Autenticación EPA falló (HTTP {auth.status_code}).')
    headers = {'Authorization': 'Bearer ' + auth.json()['access_token']}
    base = f'https://graph.microsoft.com/v1.0/sites/{site}/drive'
    url = f'{base}/root:/{quote(folder, safe="/")}:/children?$top=200'
    items = []
    while url:
        response = requests.get(url, headers=headers, timeout=60)
        if response.status_code != 200:
            raise ValueError(f'No se pudo listar EPA_CALIDAD (HTTP {response.status_code}).')
        page = response.json()
        items.extend(page.get('value', []))
        url = page.get('@odata.nextLink')
        if url and not url.startswith('https://graph.microsoft.com/'):
            raise ValueError('Paginación Graph inesperada.')
    files = [i for i in items if 'file' in i and re.fullmatch(r'EPA[ _]PERSONA[ _]IPSO_.*\.xlsx', i['name'], re.I)]
    if not files:
        raise ValueError('EPA_CALIDAD no contiene bases de encuestas.')
    for item in files:
        name = item['name']
        if Path(name).name != name or '/' in name or '\\' in name:
            raise ValueError('Nombre de archivo EPA no válido.')
        response = requests.get(f'{base}/items/{quote(item["id"], safe="")}/content', headers=headers, timeout=(20,120))
        if response.status_code != 200:
            raise ValueError(f'No se pudo descargar una base EPA (HTTP {response.status_code}).')
        (destination / name).write_bytes(response.content)
    print(f'Descargadas {len(files)} bases EPA desde SharePoint.')

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', type=Path, default=ROOT.parent.parent / 'BASES ENTEL' / 'EPA_CALIDAD')
    parser.add_argument('--html', type=Path, default=ROOT / 'PanelCallbackEpa.html')
    parser.add_argument('--graph', action='store_true', help='Descargar todas las bases desde SharePoint con los secrets de Azure')
    args = parser.parse_args()
    if args.graph:
        with tempfile.TemporaryDirectory(prefix='epa-') as folder:
            download_graph(Path(folder))
            data = build(Path(folder))
    else:
        data = build(args.source)
    update_html(args.html, data)
    print(f'Actualizadas {len(data["registros"])} encuestas de {len(data["fuentes"])} archivos.')

if __name__ == '__main__':
    try:
        main()
    except Exception as exc:
        # Evitar URLs de descarga firmadas o respuestas de autenticación en logs.
        import sys
        print(f'ERROR EPA: {exc if isinstance(exc, ValueError) else type(exc).__name__}', file=sys.stderr)
        sys.exit(1)
